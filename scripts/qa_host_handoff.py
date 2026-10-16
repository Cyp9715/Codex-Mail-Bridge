"""Standalone host QA; document libraries stay outside the plugin runtime."""
from email import policy
from email.message import EmailMessage
from contextlib import closing
from io import BytesIO
from pathlib import Path
import hashlib
import json
import os
import sqlite3
import subprocess
import sys
import tempfile

from docx import Document
from pypdf import PdfReader

root = Path(__file__).resolve().parents[1]
plugin = root / "plugins/codex-mail-bridge"
sys.path[:0] = [str(plugin / "server"), str(plugin / "tests")]
import accounts
from attachment_demo import pdf, setup
from create_demo import create_demo

(root / "work").mkdir(exist_ok=True)
with tempfile.TemporaryDirectory(dir=root / "work", prefix="synthetic-handoff-") as temporary:
    environment = create_demo(temporary)
    environment["MAIL_BRIDGE_EXPORT_ROOT"] = str(Path(temporary) / "exports")
    os.environ.update(environment)
    profile = accounts.profiles()[0]
    accounts.select_account(profile["id"], "account1", confirm=True)
    folder, _, _ = setup(profile["path"])
    document = Document()
    document.add_paragraph("Synthetic original-file handoff.")
    stream = BytesIO()
    document.save(stream)
    originals = {"agenda.pdf": pdf(), "handoff.docx": stream.getvalue()}
    message = EmailMessage(policy=policy.SMTP)
    message["Message-ID"] = "<demo-101@example.test>"
    message["From"] = "alex@example.test"
    message["To"] = "demo@example.test"
    message["Subject"] = "Project Aurora meeting"
    message.set_content("Synthetic body.")
    for name, data in originals.items():
        subtype = "pdf" if name.endswith(".pdf") else "vnd.openxmlformats-officedocument.wordprocessingml.document"
        message.add_attachment(data, maintype="application", subtype=subtype, filename=name)
    folder.write_bytes(b"From - synthetic date\r\n" + message.as_bytes())
    with closing(sqlite3.connect(profile["path"] / "global-messages-db.sqlite")) as db:
        db.execute("UPDATE messagesText_content SET c2attachmentNames=? WHERE docid=101", ("agenda.pdf handoff.docx",))
        db.commit()
    launcher = plugin / "scripts/mail_bridge.cjs"
    paths = []
    exported = []
    try:
        def command(*arguments):
            process = subprocess.run(["node", str(launcher), *arguments], capture_output=True,
                                     text=True, encoding="utf-8", timeout=20,
                                     env={**os.environ, "MAIL_BRIDGE_PYTHON": sys.executable})
            assert process.returncode == 0 and process.stderr == "", "Synthetic CLI failed"
            result = json.loads(process.stdout)
            assert result["ok"]
            return result["result"]

        def tool(name, arguments):
            argument_path = Path(temporary) / "private-query.json"
            argument_path.write_text(json.dumps(arguments), encoding="utf-8")
            return command("--call", name, "--arguments-file", str(argument_path))

        assert command("--version")["version"] == "0.7.0"
        found = tool("thunderbird_search_attachments", {"filename": "handoff.docx", "after": "2026-10-01T00:00:00Z"})
        assert [m["id"] for m in found["messages"]] == [101]
        listing = tool("thunderbird_list_attachments", {"message_id": 101})
        for attachment in listing["attachments"]:
            meta = tool("thunderbird_get_attachment", {"message_id": 101, "attachment_id": attachment["id"]})
            path = Path(meta["local_path"])
            paths.append(path)
            exported.append(meta["export_id"])
            assert path.read_bytes() == originals[attachment["name"]]
            assert meta["sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
            assert meta["document_parsed"] is False
            # This is the separate host/file-tool process, not plugin parsing.
            if attachment["name"].endswith(".pdf"):
                assert "Synthetic agenda" in PdfReader(path).pages[0].extract_text()
            else:
                assert Document(path).paragraphs[0].text == "Synthetic original-file handoff."
        assert all(path.exists() for path in paths), "Exports must survive CLI exit for host reading"
        for export_id in exported:
            assert command("--release", export_id)["released"]
        assert all(not path.exists() for path in paths)
    finally:
        for export_id in exported:
            command("--release", export_id)

print(json.dumps({"files": 2, "original_bytes_and_hashes_verified": True,
                  "separate_host_pdf_and_docx_read": True, "survives_cli_exit": True,
                  "explicit_release_verified": True, "mcp_server_required": False,
                  "installed_plugin_session_verified": False}))
